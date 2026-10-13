# Fills the Excel template with random test days and inventory counts, then recalculates it
# with LibreOffice and writes the Excel's own results to test/out/ for excel-parity.mjs.
import json, os, random, subprocess, sys
import openpyxl
src = sys.argv[1]
out = os.path.join(os.path.dirname(__file__), 'out'); os.makedirs(out, exist_ok=True)
random.seed(7)
wb = openpyxl.load_workbook(src); ws = wb['Daily Entry']; days = []
for row in range(9, 68):
    if random.random() < 0.35: continue
    v = dict(C=random.randint(10, 80), D=round(random.uniform(0, 300), 2), E=round(random.uniform(0, 120), 2), F=0, G=round(random.uniform(0, 50), 2),
             H=round(random.uniform(0, 30), 2), I=round(random.uniform(0, 150), 2), J=round(random.uniform(0, 60), 2), K=round(random.uniform(0, 40), 2),
             L=round(random.uniform(0, 30), 2), M=round(random.uniform(100, 500), 2), N=round(random.uniform(300, 2500), 2), O=round(random.uniform(100, 900), 2))
    for col, val in v.items(): ws[f'{col}{row}'] = val
    days.append({'row': row, **v})
inv = {'January': dict(F12=1200, H12=800, F13=1500, H13=700), 'February': dict(F12=1500, H12=700, F13=1100, H13=900)}
for m, cells in inv.items():
    for c, val in cells.items(): wb[m][c] = val
filled = os.path.join(out, 'filled.xlsx'); wb.save(filled)
json.dump({'days': days, 'inv': inv}, open(os.path.join(out, 'input.json'), 'w'))
subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', os.path.join(out, 'recalced'), filled], check=True, capture_output=True)
wb = openpyxl.load_workbook(os.path.join(out, 'recalced', 'filled.xlsx'), data_only=True)
labels = {16: 'covers', 17: 'revenue', 18: 'foodCogs', 19: 'bevCogs', 20: 'foodCostPct', 21: 'bevCostPct', 22: 'costOfSalesPct', 23: 'laborAll', 24: 'laborPct', 25: 'primePct', 26: 'avgSpend', 27: 'foodCostPerCover'}
res = {m: {v: wb[m][f'G{r}'].value for r, v in labels.items()} for m in ['January', 'February']}
c = wb['Cummulatief']
res['YTD'] = dict(covers=c['C17'].value, revenue=c['D17'].value, foodCogs=c['E17'].value, foodCostPct=c['F17'].value, bevCogs=c['G17'].value, bevCostPct=c['H17'].value,
                  costOfSalesPct=c['I17'].value, laborAll=c['J17'].value, laborPct=c['K17'].value, primePct=c['L17'].value, avgSpend=c['M17'].value, foodCostPerCover=c['N17'].value)
res['nonfood'] = {'January': c['C22'].value, 'February': c['C23'].value}
json.dump(res, open(os.path.join(out, 'excel_out.json'), 'w'))
print(len(days), 'test days written and recalculated')
